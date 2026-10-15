#!/usr/bin/env python3
"""tracker.py - read and write job-tracker.xlsx, the one live log of the job hunt.

The tracker is one workbook with one sheet named "Applications". Row 1 is the
header and stays frozen. The columns, in this exact order:

  Status | Company | Role | Location | Band | Next action | Package | Notes

Rules this script enforces (see references/tracker-contract.md):

  * Every write backs the file up first, to
    job-tracker.backup-<YYYY-MM-DD>-<tag>.xlsx in the same folder.
  * Notes are newest first. Each entry starts with a date and entries are
    joined with " || ". Old notes are never overwritten:
        14 Oct 2026: Phone screen booked || 2 Oct 2026: Applied
  * Status uses a fixed vocabulary so filters work. A short detail may follow a
    dash, e.g. "Interviewing - round 2 on 15 Oct".
  * New rows copy the cell styles of the row above.
  * update --match must find exactly one row, or the script stops (exit 2) and
    prints the candidates.

Commands (all take --file, default job-tracker.xlsx, and --hunt-dir, default .):

  tracker.py init
  tracker.py add --company "Acme" --role "Product Designer" --status Shortlisted \\
      [--location "Remote (US)"] [--band "$150,000 - $180,000"] [--next "Read the JD"] \\
      [--package applications/acme-product-designer/] [--note "Found in the sweep"]
  tracker.py update --match "Acme" [--role "Designer"] [--status "Applied"] \\
      [--next "Wait for reply"] [--package ...] [--note "Applied on the site"]
  tracker.py list [--status Applied] [--company Acme] [--notes]
  tracker.py report

No openpyxl: the same commands work on job-tracker.csv (same columns, same
note format) and the script prints one line on how to install openpyxl.
If the hunt folder has job-tracker.csv and no job-tracker.xlsx, every command
uses the csv even when openpyxl is installed, so there is never a second log.
With no openpyxl and an existing job-tracker.xlsx, the script stops instead of
starting a csv next to it.

Exit codes: 0 = done, 2 = a problem you can fix (message says what).
"""
from __future__ import annotations

import argparse
import copy
import csv
import logging
import re
import shutil
import textwrap
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

log = logging.getLogger(__name__)

SHEET_NAME = "Applications"
DEFAULT_FILE = "job-tracker.xlsx"

# Locked column order. Do not reorder without changing the contract.
COLUMNS: List[str] = [
    "Status",
    "Company",
    "Role",
    "Location",
    "Band",
    "Next action",
    "Package",
    "Notes",
]

COLUMN_WIDTHS: Dict[str, int] = {
    "Status": 22,
    "Company": 20,
    "Role": 40,
    "Location": 26,
    "Band": 24,
    "Next action": 48,
    "Package": 30,
    "Notes": 80,
}

# The status vocabulary, in the order the report prints it.
STATUSES: List[str] = [
    "Shortlisted",
    "Ready to apply",
    "Applied",
    "Applied - confirmed",
    "Interviewing",
    "Offer",
    "Rejected",
    "Closed",
    "Withdrawn",
    "Recruiter contact",
    "Inbound",
    "Parked",
]

# Status cell colours. Blue = not sent yet, green = sent, yellow = live
# conversation, orange = someone reached out, grey = finished or parked.
STATUS_FILLS: Dict[str, str] = {
    "Shortlisted": "EAF1FB",
    "Ready to apply": "D6E4F7",
    "Applied": "C6E7C8",
    "Applied - confirmed": "C6E7C8",
    "Interviewing": "FFF2A8",
    "Offer": "FFE08A",
    "Rejected": "E9E9E9",
    "Closed": "E9E9E9",
    "Withdrawn": "E9E9E9",
    "Recruiter contact": "FDE9C9",
    "Inbound": "FDE9C9",
    "Parked": "E9E9E9",
}

# Rows with these statuses always show up under "Needs you" in the report.
NEEDS_YOU_STATUSES = {"Interviewing", "Offer", "Inbound", "Recruiter contact"}
NEEDS_YOU_PATTERN = re.compile(r"\b(?:you|your call|needs you)\b", re.IGNORECASE)

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
NOTE_SEPARATOR = " || "

HEADER_FILL = "1F2933"
HEADER_FONT_COLOR = "FFFFFF"
BORDER_COLOR = "D8DEE4"

INSTALL_LINE = ("Note: openpyxl is not installed, so I used {csv_name} instead of the "
                "spreadsheet. For the styled .xlsx, run: pip install openpyxl")


class TrackerError(Exception):
    """A problem the user can fix. The message is printed as is; exit code 2."""


@dataclass
class TrackerRow:
    """One tracker row. `number` is the sheet row (header is row 1)."""

    number: int
    status: str
    company: str
    role: str
    location: str
    band: str
    next_action: str
    package: str
    notes: str

    def values(self) -> List[str]:
        """Return the cell values in column order."""
        return [self.status, self.company, self.role, self.location,
                self.band, self.next_action, self.package, self.notes]

    @classmethod
    def from_values(cls, number: int, values: List[Optional[object]]) -> "TrackerRow":
        """Build a row from raw cell values, padding missing cells with ''."""
        cells = [_clean(v) for v in values[:len(COLUMNS)]]
        cells += [""] * (len(COLUMNS) - len(cells))
        return cls(number, *cells)

    def is_empty(self) -> bool:
        """True when every cell is blank."""
        return not any(v.strip() for v in self.values())


@dataclass
class Store:
    """Where the tracker lives and which backend reads it."""

    path: Path
    use_xlsx: bool
    note: Optional[str] = None
    xlsx_twin: Optional[Path] = None
    csv_is_the_log: bool = False

    def check_no_hidden_xlsx(self) -> None:
        """Stop if a real .xlsx exists but we can only use the csv (would split the log)."""
        if (not self.use_xlsx and self.xlsx_twin is not None
                and self.xlsx_twin.exists() and not self.path.exists()):
            raise TrackerError(
                f"{self.xlsx_twin.name} exists, but openpyxl is not installed, so I cannot "
                "read it. Run: pip install openpyxl")


# ---------------------------------------------------------------- helpers

def _clean(value: Optional[object]) -> str:
    """Turn a cell value into a plain string."""
    if value is None:
        return ""
    return str(value).strip()


def openpyxl_available() -> bool:
    """True if openpyxl can be imported."""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def format_note_date(day: date) -> str:
    """Format a date the way notes start: '14 Oct 2026' (no leading zero)."""
    return f"{day.day} {MONTHS[day.month - 1]} {day.year}"


def parse_day(text: Optional[str]) -> date:
    """Parse --date (YYYY-MM-DD). None means today."""
    if not text:
        return date.today()
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        raise TrackerError(f"--date must look like 2026-10-14, not '{text}'.")


def prepend_note(old_notes: str, note: str, day: date) -> str:
    """Put a dated note in front of the old notes. Never drops old notes."""
    entry = f"{format_note_date(day)}: {note.strip()}"
    old = old_notes.strip()
    return f"{entry}{NOTE_SEPARATOR}{old}" if old else entry


def base_status(status: str) -> Optional[str]:
    """Return the vocabulary word a status starts with, or None.

    The longest match wins, so "Applied - confirmed" is not read as "Applied".
    A detail may follow after " - ", e.g. "Interviewing - round 2".
    """
    lowered = status.strip().lower()
    for word in sorted(STATUSES, key=len, reverse=True):
        w = word.lower()
        if lowered == w or lowered.startswith(w + " - "):
            return word
    return None


def normalize_status(status: str) -> str:
    """Check a status against the vocabulary and fix the case of its first part."""
    word = base_status(status)
    if word is None:
        allowed = ", ".join(STATUSES)
        raise TrackerError(
            f"'{status}' is not a tracker status. Use one of: {allowed}. "
            "You may add a detail after a dash, like 'Interviewing - round 2'.")
    return word + status.strip()[len(word):]


def needs_you(row: TrackerRow) -> bool:
    """True if the row needs the user: a live status or a next action naming them."""
    if base_status(row.status) in NEEDS_YOU_STATUSES:
        return True
    return bool(NEEDS_YOU_PATTERN.search(row.next_action))


def resolve_store(hunt_dir: Path, file_arg: str, no_openpyxl: bool = False) -> Store:
    """Work out the tracker path and backend (xlsx or csv)."""
    path = Path(file_arg)
    if not path.is_absolute():
        path = hunt_dir / path
    if path.suffix.lower() == ".csv":
        return Store(path=path, use_xlsx=False)
    if no_openpyxl or not openpyxl_available():
        csv_path = path.with_suffix(".csv")
        return Store(path=csv_path, use_xlsx=False,
                     note=INSTALL_LINE.format(csv_name=csv_path.name), xlsx_twin=path)
    csv_path = path.with_suffix(".csv")
    if not path.exists() and csv_path.exists():
        return Store(path=csv_path, use_xlsx=False, csv_is_the_log=True)
    return Store(path=path, use_xlsx=True)


def backup(path: Path, tag: str) -> Optional[Path]:
    """Copy the tracker to <stem>.backup-<YYYY-MM-DD>-<tag><suffix> before a write.

    If that name is taken (two writes the same day with the same tag), a counter
    is added so an earlier backup is never overwritten.
    """
    if not path.exists():
        return None
    safe_tag = re.sub(r"[^A-Za-z0-9_-]+", "-", tag).strip("-") or "write"
    stem = f"{path.stem}.backup-{date.today().isoformat()}-{safe_tag}"
    target = path.with_name(stem + path.suffix)
    counter = 2
    while target.exists():
        target = path.with_name(f"{stem}-{counter}{path.suffix}")
        counter += 1
    shutil.copy2(path, target)
    log.debug("backed up %s to %s", path, target)
    return target


def find_matches(rows: List[TrackerRow], company: str,
                 role: Optional[str]) -> List[TrackerRow]:
    """Rows whose Company contains `company` (and Role contains `role`), any case."""
    c = company.strip().lower()
    r = role.strip().lower() if role else None
    hits = [row for row in rows if c in row.company.lower()]
    if r:
        hits = [row for row in hits if r in row.role.lower()]
    return hits


def _truncate(text: str, width: int) -> str:
    """Cut text to width characters, ending with '...' if it was cut."""
    text = " ".join(text.split())
    if len(text) <= width:
        return text
    return text[: max(width - 3, 0)] + "..."


def format_table(rows: List[TrackerRow], notes: bool = False) -> str:
    """A compact plain-text table for the terminal. notes=True adds each row's full Notes."""
    spec: List[Tuple[str, int]] = [("#", 4), ("Status", 30), ("Company", 22),
                                   ("Role", 36), ("Next action", 40)]
    lines = ["  ".join(name.ljust(w) for name, w in spec).rstrip()]
    lines.append("  ".join("-" * w for _, w in spec))
    for row in rows:
        cells = [str(row.number), row.status, row.company, row.role, row.next_action]
        lines.append("  ".join(_truncate(v, w).ljust(w)
                               for v, (_, w) in zip(cells, spec)).rstrip())
        if notes:
            lines.append(textwrap.fill(row.notes or "(none)", width=110,
                                       initial_indent="      Notes: ",
                                       subsequent_indent=" " * 13,
                                       break_long_words=False, break_on_hyphens=False))
    return "\n".join(lines)


# ---------------------------------------------------------------- reading

def _read_xlsx(path: Path) -> List[TrackerRow]:
    """Read rows from the Applications sheet."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise TrackerError(f"{path.name} has no sheet named '{SHEET_NAME}'.")
        ws = wb[SHEET_NAME]
        rows: List[TrackerRow] = []
        for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
            if number == 1:
                _check_header([_clean(v) for v in values], path)
                continue
            row = TrackerRow.from_values(number, list(values))
            if not row.is_empty():
                rows.append(row)
        return rows
    finally:
        wb.close()


def _read_csv(path: Path) -> List[TrackerRow]:
    """Read rows from the csv fallback."""
    rows: List[TrackerRow] = []
    with path.open("r", encoding="utf-8", newline="") as f:
        for number, values in enumerate(csv.reader(f), start=1):
            if number == 1:
                _check_header([_clean(v) for v in values], path)
                continue
            row = TrackerRow.from_values(number, list(values))
            if not row.is_empty():
                rows.append(row)
    return rows


def _check_header(header: List[str], path: Path) -> None:
    """Stop if row 1 is not the tracker header (protects other spreadsheets)."""
    if header[:len(COLUMNS)] != COLUMNS:
        raise TrackerError(
            f"Row 1 of {path.name} is not the tracker header. "
            f"Expected: {', '.join(COLUMNS)}.")


def read_rows(path: Path) -> List[TrackerRow]:
    """Read every non-empty tracker row from an .xlsx or .csv file.

    Other scripts (sweep_boards.py) use this. If the .xlsx is missing, or
    openpyxl is not installed, it reads the .csv next to it if there is one.
    """
    if path.suffix.lower() == ".csv":
        return _read_csv(path)
    if not path.exists() and path.with_suffix(".csv").exists():
        return _read_csv(path.with_suffix(".csv"))
    if not openpyxl_available():
        csv_path = path.with_suffix(".csv")
        if csv_path.exists():
            return _read_csv(csv_path)
        raise TrackerError(
            f"openpyxl is not installed, so I cannot read {path.name}. "
            "Run: pip install openpyxl")
    return _read_xlsx(path)


def load(store: Store) -> List[TrackerRow]:
    """Read the tracker, or stop with a clear message if it is missing."""
    store.check_no_hidden_xlsx()
    if not store.path.exists():
        raise TrackerError(
            f"No tracker at {store.path}. Run 'tracker.py init' first.")
    return _read_xlsx(store.path) if store.use_xlsx else _read_csv(store.path)


# ---------------------------------------------------------------- xlsx writing

def _styles() -> Tuple[Any, ...]:
    """Import openpyxl style classes lazily (openpyxl is optional)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return Alignment, Border, Font, PatternFill, Side


def _apply_default_data_style(ws: Any, row: int) -> None:
    """Style a data row from scratch (used for the first row under the header)."""
    Alignment, Border, Font, PatternFill, Side = _styles()
    thin = Side(style="thin", color=BORDER_COLOR)
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(name="Calibri", size=10)
        cell.fill = PatternFill(fill_type=None)


def _paint_status(ws: Any, row: int, status: str) -> None:
    """Colour the Status cell for its status and make it bold."""
    Alignment, Border, Font, PatternFill, Side = _styles()
    cell = ws.cell(row=row, column=1)
    word = base_status(status)
    colour = STATUS_FILLS.get(word or "")
    cell.fill = PatternFill("solid", fgColor=colour) if colour else PatternFill(fill_type=None)
    cell.font = Font(name="Calibri", size=10, bold=True)


def _last_used_row(ws: Any) -> int:
    """The last row with any value (formatting alone does not count)."""
    for row in range(ws.max_row, 1, -1):
        if any(ws.cell(row=row, column=c).value not in (None, "")
               for c in range(1, len(COLUMNS) + 1)):
            return row
    return 1


def _open_sheet(path: Path) -> Tuple[Any, Any]:
    """Open the workbook for writing and return (workbook, sheet)."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    if SHEET_NAME not in wb.sheetnames:
        raise TrackerError(f"{path.name} has no sheet named '{SHEET_NAME}'.")
    ws = wb[SHEET_NAME]
    _check_header([_clean(ws.cell(row=1, column=c).value)
                   for c in range(1, len(COLUMNS) + 1)], path)
    return wb, ws


def _set_filter(ws: Any, last_row: int) -> None:
    """Keep the header filter covering every row."""
    from openpyxl.utils import get_column_letter

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(last_row, 1)}"


def xlsx_init(path: Path) -> None:
    """Create a new tracker workbook with the styled, frozen header."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    Alignment, Border, Font, PatternFill, Side = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Calibri", size=11, bold=True, color=HEADER_FONT_COLOR)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[name]
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    _set_filter(ws, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def xlsx_add(path: Path, row: TrackerRow) -> int:
    """Append a row, copying the styles of the row above. Returns the row number."""
    wb, ws = _open_sheet(path)
    above = _last_used_row(ws)
    new = above + 1
    for col, value in enumerate(row.values(), start=1):
        ws.cell(row=new, column=col).value = value
    if above >= 2:
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=new, column=col)._style = copy.copy(
                ws.cell(row=above, column=col)._style)
    else:
        _apply_default_data_style(ws, new)
    _paint_status(ws, new, row.status)
    _set_filter(ws, new)
    wb.save(path)
    return new


def xlsx_update(path: Path, row: TrackerRow) -> None:
    """Write a changed row back to its sheet row."""
    wb, ws = _open_sheet(path)
    old_status = _clean(ws.cell(row=row.number, column=1).value)
    for col, value in enumerate(row.values(), start=1):
        ws.cell(row=row.number, column=col).value = value
    if old_status != row.status:
        _paint_status(ws, row.number, row.status)
    wb.save(path)


# ---------------------------------------------------------------- csv writing

def csv_write_all(path: Path, rows: List[TrackerRow]) -> None:
    """Rewrite the whole csv (header + rows)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow(row.values())


# ---------------------------------------------------------------- commands

def cmd_init(store: Store, args: argparse.Namespace) -> int:
    """Create the tracker if it does not exist yet."""
    store.check_no_hidden_xlsx()
    if store.csv_is_the_log:
        print(f"{store.path.name} is already the tracker in this folder, so I did not start "
              f"a {store.path.with_suffix('.xlsx').name}. Left it as it is.")
        return 0
    if store.path.exists():
        print(f"Tracker already exists at {store.path}. Left it as it is.")
        return 0
    if store.use_xlsx:
        xlsx_init(store.path)
    else:
        csv_write_all(store.path, [])
    print(f"Created {store.path} with the {len(COLUMNS)} tracker columns.")
    return 0


def cmd_add(store: Store, args: argparse.Namespace) -> int:
    """Add one row after backing up."""
    rows = load(store)
    status = normalize_status(args.status)
    company = args.company.strip()
    role = args.role.strip()
    for existing in rows:
        if (existing.company.lower() == company.lower()
                and existing.role.lower() == role.lower()):
            raise TrackerError(
                f"Row {existing.number} already tracks {existing.company} / "
                f"{existing.role} ({existing.status}). Use 'tracker.py update' instead.")
    notes = prepend_note("", args.note, parse_day(args.date)) if args.note else ""
    new = TrackerRow(
        number=0, status=status, company=company, role=role,
        location=(args.location or "").strip(), band=(args.band or "").strip(),
        next_action=(args.next or "").strip(), package=(args.package or "-").strip() or "-",
        notes=notes)
    saved = backup(store.path, args.tag or "add")
    if store.use_xlsx:
        new.number = xlsx_add(store.path, new)
    else:
        new.number = (rows[-1].number + 1) if rows else 2
        csv_write_all(store.path, rows + [new])
    print(f"Added row {new.number}: {new.status} | {new.company} | {new.role}")
    if saved:
        print(f"Backup: {saved}")
    return 0


def cmd_update(store: Store, args: argparse.Namespace) -> int:
    """Update one row (status, next action, package, prepend a note) after backing up."""
    if not any([args.status, args.next, args.package, args.note]):
        raise TrackerError(
            "Nothing to change. Give at least one of --status, --next, --package, --note.")
    rows = load(store)
    hits = find_matches(rows, args.match, args.role)
    if len(hits) != 1:
        what = f"'{args.match}'" + (f" with role '{args.role}'" if args.role else "")
        if not hits:
            raise TrackerError(f"No row matches company {what}. Nothing changed.")
        print(f"{len(hits)} rows match company {what}. Nothing changed. "
              "Add --role to pick one:", file=sys.stderr)
        print(format_table(hits), file=sys.stderr)
        return 2
    row = hits[0]
    if args.status:
        row.status = normalize_status(args.status)
    if args.next:
        row.next_action = args.next.strip()
    if args.package:
        row.package = args.package.strip()
    if args.note:
        row.notes = prepend_note(row.notes, args.note, parse_day(args.date))
    saved = backup(store.path, args.tag or "update")
    if store.use_xlsx:
        xlsx_update(store.path, row)
    else:
        csv_write_all(store.path, rows)
    print(f"Updated row {row.number}: {row.status} | {row.company} | {row.role}")
    if args.note:
        print(f"Notes now start: {_truncate(row.notes, 120)}")
    if saved:
        print(f"Backup: {saved}")
    return 0


def cmd_list(store: Store, args: argparse.Namespace) -> int:
    """Print rows as a compact table, filtered by status and company."""
    rows = load(store)
    if args.status:
        want = args.status.strip().lower()
        rows = [r for r in rows if r.status.lower() == want
                or r.status.lower().startswith(want)]
    if args.company:
        want = args.company.strip().lower()
        rows = [r for r in rows if want in r.company.lower()]
    if not rows:
        print("No rows match.")
        return 0
    print(format_table(rows, notes=args.notes))
    print(f"\n{len(rows)} row(s).")
    return 0


def cmd_report(store: Store, args: argparse.Namespace) -> int:
    """Print counts by status and the rows that need the user."""
    rows = load(store)
    counts: Dict[str, int] = {}
    for row in rows:
        key = base_status(row.status) or (row.status or "(blank)")
        counts[key] = counts.get(key, 0) + 1
    order = [s for s in STATUSES if s in counts] + sorted(
        k for k in counts if k not in STATUSES)
    print(f"Tracker: {store.path} ({len(rows)} rows)\n")
    print("By status:")
    for key in order:
        print(f"  {key.ljust(22)}{counts[key]}")
    flagged = [r for r in rows if needs_you(r)]
    print(f"\nNeeds you ({len(flagged)}):")
    if not flagged:
        print("  Nothing right now.")
    for row in flagged:
        nxt = row.next_action or "(no next action written)"
        print(f"  Row {row.number}: {row.status} | {row.company} | {row.role}")
        print(f"      Next: {nxt}")
    return 0


# ---------------------------------------------------------------- main

EXAMPLES = """examples:
  tracker.py init --hunt-dir ~/job-hunt
  tracker.py add --company "Acme" --role "Product Designer" --status Shortlisted --note "Found in the sweep"
  tracker.py update --match "Acme" --status "Applied" --next "Wait for a reply" --note "Applied on the site"
  tracker.py list --status Applied
  tracker.py list --company Acme --notes
  tracker.py report

statuses: """ + ", ".join(STATUSES) + "\n"


def build_parser() -> argparse.ArgumentParser:
    """Build the argparse parser with the five subcommands."""
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--file", default=DEFAULT_FILE,
                        help="tracker file name or path (default: job-tracker.xlsx)")
    common.add_argument("--hunt-dir", default=".", type=Path,
                        help="the hunt folder; --file is read relative to it (default: .)")
    common.add_argument("--no-openpyxl", action="store_true", help=argparse.SUPPRESS)

    write = argparse.ArgumentParser(add_help=False)
    write.add_argument("--tag", help="word for the backup file name (default: the command)")
    write.add_argument("--date", help="date for the note, YYYY-MM-DD (default: today)")

    parser = argparse.ArgumentParser(
        description="Read and write job-tracker.xlsx, the one live log of the hunt. "
                    "Every write backs the file up first.",
        epilog=EXAMPLES,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", metavar="command")
    sub.required = True

    sub.add_parser("init", parents=[common],
                   help="create the tracker with the styled, frozen header")

    p_add = sub.add_parser("add", parents=[common, write], help="add one row")
    p_add.add_argument("--company", required=True)
    p_add.add_argument("--role", required=True)
    p_add.add_argument("--status", required=True,
                       help="one of: " + ", ".join(STATUSES))
    p_add.add_argument("--location")
    p_add.add_argument("--band", help="published pay band, or blank")
    p_add.add_argument("--next", help="the next action, in plain words")
    p_add.add_argument("--package", help="role folder, e.g. applications/acme-designer/ (default: -)")
    p_add.add_argument("--note", help="first note; the date is added for you")

    p_up = sub.add_parser("update", parents=[common, write],
                          help="change one row; --match must find exactly one")
    p_up.add_argument("--match", required=True, help="part of the company name")
    p_up.add_argument("--role", help="part of the role title, to pick one row")
    p_up.add_argument("--status", help="new status")
    p_up.add_argument("--next", help="new next action")
    p_up.add_argument("--package", help="new role folder")
    p_up.add_argument("--note", help="note to put in front of the old notes")

    p_list = sub.add_parser("list", parents=[common], help="print rows as a table")
    p_list.add_argument("--status", help="only rows whose status starts with this")
    p_list.add_argument("--company", help="only rows whose company contains this")
    p_list.add_argument("--notes", action="store_true",
                        help="also print each row's full Notes (newest first) under it")

    sub.add_parser("report", parents=[common],
                   help="counts by status and the rows that need you")
    return parser


COMMANDS = {
    "init": cmd_init,
    "add": cmd_add,
    "update": cmd_update,
    "list": cmd_list,
    "report": cmd_report,
}


def main(argv: List[str]) -> int:
    """Entry point. Returns the exit code."""
    logging.basicConfig(level=logging.WARNING, format="%(message)s")
    args = build_parser().parse_args(argv)
    store = resolve_store(args.hunt_dir, args.file, args.no_openpyxl)
    try:
        store.check_no_hidden_xlsx()
        if store.note:
            print(store.note, file=sys.stderr)
        return COMMANDS[args.command](store, args)
    except TrackerError as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
